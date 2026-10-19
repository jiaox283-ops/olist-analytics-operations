"""模块六 第 5 步：把 results 下的 CSV 汇总成一个可直接翻阅的 Excel。"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parents[1]
RES = BASE / 'results' / 'module6'
OUT_XLSX = RES / 'module6_summary.xlsx'

SHEETS = [
    ('01_复现校验', '01_score_reproduction_check.json', 'json'),
    ('02_策略比较', '03_strategy_comparison.csv', 'csv'),
    ('03_决策摘要', '07_decision_summary.csv', 'csv'),
    ('04_六情景敏感性', '06_sensitivity_grid.csv', 'csv'),
    ('05_盈亏平衡', '04_break_even.csv', 'csv'),
    ('06_价值排序', '09_value_targeting.csv', 'csv'),
    ('07_AB功效', '12_ab_power.csv', 'csv'),
    ('08_AB的MDE', '13_ab_mde.csv', 'csv'),
    ('09_CUPED', '14_ab_cuped.csv', 'csv'),
    ('10_假设参数', '08_assumptions.json', 'json'),
    ('11_风险分级', '16_risk_tiers.csv', 'csv'),
]


def json_to_df(path: Path) -> pd.DataFrame:
    obj = json.loads(path.read_text(encoding='utf-8'))

    def walk(o, prefix=''):
        if isinstance(o, dict):
            for k, v in o.items():
                yield from walk(v, f'{prefix}{k}.')
        else:
            yield prefix.rstrip('.'), o

    rows = list(walk(obj))
    return pd.DataFrame(rows, columns=['参数', '取值'])


def main() -> None:
    with pd.ExcelWriter(OUT_XLSX, engine='openpyxl') as w:
        intro = pd.DataFrame([
            ['项目', '电商平台经营分析、用户价值分层与物流履约风险预测'],
            ['模块', '模块六：运营策略与效果评估'],
            ['数据', 'Olist Brazilian E-Commerce Public Dataset（模块一清洗后的 mart_order）'],
            ['测试集', '2018-06-01 至 2018-08-29 已送达且延迟标签有效的 18,603 笔订单；90 个有订单日'],
            ['金额单位', '巴西雷亚尔 BRL'],
            ['重要口径 1', '所有金额均为未经验证参数下的历史标签情景值，不是观测到的干预收益。'],
            ['重要口径 2', '06_价值排序中的标签情景值与分数情景值均是假设；score_to_label_ratio 仅用于诊断分数偏差。'],
            ['重要口径 3', '本模块未上线，没有观测到真实干预收益。'],
            ['重要口径 4', '风险分级是测试集内分位，不是可直接上线的固定阈值；Logistic 为测试后探索性比较。'],
            ['前置依赖', '模块四冻结模型与分数；模块六仅做策略层与决策层分析，未重新调参。'],
        ], columns=['说明', '内容'])
        intro.to_excel(w, sheet_name='00_说明', index=False)

        for sheet, fname, kind in SHEETS:
            p = RES / fname
            if not p.exists():
                continue
            df = json_to_df(p) if kind == 'json' else pd.read_csv(p)
            df = df.astype(object).where(pd.notna(df), None)
            df.to_excel(w, sheet_name=sheet, index=False)

    # 轻量格式化：表头加粗+底色，冻结首行，列宽自适应
    from openpyxl import load_workbook
    wb = load_workbook(OUT_XLSX)
    head_fill = PatternFill('solid', fgColor='DCE6F1')
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.freeze_panes = 'A2'
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)
    wb.save(OUT_XLSX)
    print('已写出', OUT_XLSX)


if __name__ == '__main__':
    main()
